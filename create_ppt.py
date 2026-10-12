from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Brand colors
DARK_BG = RGBColor(0x0F, 0x17, 0x2A)       # Deep navy
ACCENT_BLUE = RGBColor(0x00, 0x96, 0xD6)   # StemCyte blue
ACCENT_TEAL = RGBColor(0x00, 0xB4, 0xD8)   # Teal accent
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xB0, 0xC4, 0xDE)
CARD_BG = RGBColor(0x1A, 0x23, 0x3B)       # Card background
GREEN = RGBColor(0x00, 0xC8, 0x53)
ORANGE = RGBColor(0xFF, 0x8C, 0x00)
RED = RGBColor(0xFF, 0x45, 0x45)

def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color, border_color=None, corner_radius=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.fill.solid()
        shape.line.fill.fore_color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=18, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_bullet_text(slide, left, top, width, height, items, font_size=14, color=WHITE, title=None, title_size=16, title_color=ACCENT_BLUE):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(title_size)
        p.font.color.rgb = title_color
        p.font.bold = True
        p.font.name = 'Calibri'
        p.space_after = Pt(8)
        start_new = True
    else:
        start_new = False

    for i, item in enumerate(items):
        if start_new or i > 0:
            p = tf.add_paragraph()
        else:
            p = tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = 'Calibri'
        p.space_after = Pt(4)
        p.level = 0
    return txBox

def add_accent_line(slide, left, top, width, color=ACCENT_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

# ============================================================
# SLIDE 1: Title
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
set_slide_bg(slide, DARK_BG)

# Accent line at top
add_accent_line(slide, Inches(0), Inches(0), Inches(13.333), ACCENT_BLUE)

# Title
add_text_box(slide, Inches(1), Inches(1.5), Inches(11), Inches(1),
    "CryoSync", font_size=54, color=ACCENT_BLUE, bold=True)

add_text_box(slide, Inches(1), Inches(2.5), Inches(11), Inches(0.8),
    "Cold-Chain Intelligence & Cord Blood Banking Platform",
    font_size=28, color=WHITE, bold=False)

add_accent_line(slide, Inches(1), Inches(3.5), Inches(2), ACCENT_TEAL)

add_text_box(slide, Inches(1), Inches(4), Inches(11), Inches(0.6),
    "Designed for StemCyte India", font_size=20, color=LIGHT_GRAY)

# What we will show today
add_shape(slide, Inches(1), Inches(5.1), Inches(11.3), Inches(1.5), CARD_BG, ACCENT_TEAL)
add_text_box(slide, Inches(1.2), Inches(5.25), Inches(10.9), Inches(1.2),
    "Today's Demo: A live walkthrough of the intelligent cord blood banking platform\n"
    "-- CBU lifecycle, temperature monitoring, compliance, and AI copilot (Kyra).",
    font_size=16, color=WHITE)

add_text_box(slide, Inches(1), Inches(6.8), Inches(11), Inches(0.5),
    "AI-Powered  |  End-to-End Lifecycle Management  |  Real-Time Monitoring",
    font_size=13, color=RGBColor(0x60, 0x70, 0x90))

# ============================================================
# SLIDE 2: The Challenge + The Solution
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7),
    "The Challenge & CryoSync Solution", font_size=36, color=ACCENT_BLUE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2), ACCENT_TEAL)

# Left column: challenges
add_text_box(slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(0.4),
    "The Challenge", font_size=16, color=RED, bold=True)

challenges = [
    "Operations tracked across disconnected spreadsheets & manual logs",
    "Cold-chain visibility gaps -- temperature excursions compromise CBU viability",
    "Zoho CRM, lab reports, and databases sit in silos",
    "FACT / AABB / DCGI compliance maintained manually",
    "No single source of truth for multi-facility leadership",
]

for i, item in enumerate(challenges):
    badge = add_shape(slide, Inches(0.8), Inches(1.9 + i * 0.62), Inches(0.42), Inches(0.42), RED)
    tf = badge.text_frame
    tf.paragraphs[0].text = str(i + 1)
    tf.paragraphs[0].font.size = Pt(12)
    tf.paragraphs[0].font.color.rgb = WHITE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    add_text_box(slide, Inches(1.4), Inches(1.9 + i * 0.62), Inches(4.9), Inches(0.6),
        item, font_size=12, color=LIGHT_GRAY)

# Right column: solution cards
add_text_box(slide, Inches(7.0), Inches(1.4), Inches(5.5), Inches(0.4),
    "CryoSync Delivers", font_size=16, color=GREEN, bold=True)

solutions = [
    ("End-to-End CBU Lifecycle", "Enrollment, collection, processing, testing,\nstorage, and transplant release in one system.", ACCENT_BLUE),
    ("Real-Time Cold-Chain Alerts", "Continuous temperature monitoring with AI\nexcursion detection & auto-incidents.", ACCENT_TEAL),
    ("Unified Data Lake", "Zoho CRM + Databricks (Bronze/Silver/Gold),\nOCR lab reports, one source of truth.", ORANGE),
    ("Compliance Built-In", "Auto audit trails, deviation & incident\nmanagement for FACT / AABB / DCGI.", GREEN),
    ("AI Copilot 'Kyra'", "Natural-language answers for ops staff\n-- queries, summaries, and anomaly scans.", RGBColor(0xAA, 0x66, 0xFF)),
]

for i, (title, desc, color) in enumerate(solutions):
    top = Inches(1.9 + i * 1.06)
    add_shape(slide, Inches(7.0), top, Inches(5.5), Inches(0.95), CARD_BG, color)
    add_text_box(slide, Inches(7.2), top + Inches(0.08), Inches(5.1), Inches(0.35),
        title, font_size=13, color=color, bold=True)
    add_text_box(slide, Inches(7.2), top + Inches(0.42), Inches(5.1), Inches(0.5),
        desc, font_size=10, color=LIGHT_GRAY)

# ============================================================
# SLIDE 3: Architecture (simple & client-friendly)
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7),
    "Architecture: Simple by Design", font_size=36, color=ACCENT_BLUE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2), ACCENT_TEAL)

# --- Stage 1: Data sources ---
add_text_box(slide, Inches(0.8), Inches(1.5), Inches(3.4), Inches(0.5),
    "1  Data Sources", font_size=20, color=ACCENT_TEAL, bold=True)

source_items = [
    "Collection & operations teams",
    "Zoho CRM (customer & sales data)",
    "Lab reports & CBU records",
    "Cold-chain temperature sensors",
]
for i, item in enumerate(source_items):
    add_shape(slide, Inches(0.8), Inches(2.05 + i * 0.72), Inches(3.4), Inches(0.58), CARD_BG)
    add_text_box(slide, Inches(0.95), Inches(2.1 + i * 0.72), Inches(3.1), Inches(0.5),
        item, font_size=12, color=LIGHT_GRAY)

# Arrow 1
ar1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.45), Inches(3.1), Inches(0.8), Inches(0.75))
ar1.fill.solid()
ar1.fill.fore_color.rgb = ACCENT_TEAL
ar1.line.fill.background()

# --- Stage 2: CryoSync platform ---
add_text_box(slide, Inches(5.4), Inches(1.5), Inches(3.6), Inches(0.5),
    "2  CryoSync Platform", font_size=20, color=ACCENT_BLUE, bold=True)

platform_shape = add_shape(slide, Inches(5.4), Inches(2.05), Inches(3.6), Inches(4.0), CARD_BG, ACCENT_BLUE)
add_text_box(slide, Inches(5.6), Inches(2.2), Inches(3.2), Inches(0.9),
    "All cord blood operations in one secure web platform",
    font_size=12, color=WHITE, bold=True)

platform_items = [
    "Single source of truth -- no more manual logs or spreadsheets",
    "Real-time monitoring with automatic temperature alerts",
    "Compliance & audit trails built in",
    "AI Copilot (Kyra) -- ask questions in plain English",
]
for i, item in enumerate(platform_items):
    add_shape(slide, Inches(5.6), Inches(3.15 + i * 0.72), Inches(3.2), Inches(0.58), RGBColor(0x22, 0x2F, 0x4D))
    add_text_box(slide, Inches(5.75), Inches(3.18 + i * 0.72), Inches(3.0), Inches(0.55),
        item, font_size=11, color=LIGHT_GRAY)

# Arrow 2
ar2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(9.25), Inches(3.1), Inches(0.8), Inches(0.75))
ar2.fill.solid()
ar2.fill.fore_color.rgb = ACCENT_BLUE
ar2.line.fill.background()

# --- Stage 3: Value / outcomes ---
add_text_box(slide, Inches(10.2), Inches(1.5), Inches(3.0), Inches(0.5),
    "3  What You Get", font_size=20, color=GREEN, bold=True)

value_items = [
    "Live dashboards for leadership",
    "Instant alerts on any temperature issue",
    "Automatic compliance reports",
    "Faster, safer operational decisions",
]
for i, item in enumerate(value_items):
    add_shape(slide, Inches(10.2), Inches(2.05 + i * 0.72), Inches(2.5), Inches(0.58), CARD_BG, GREEN)
    add_text_box(slide, Inches(10.35), Inches(2.1 + i * 0.72), Inches(2.2), Inches(0.5),
        item, font_size=11, color=LIGHT_GRAY)

# Bottom summary
add_text_box(slide, Inches(0.8), Inches(6.35), Inches(11.7), Inches(0.8),
    "Every data source connects automatically into one platform -- teams work in real time,\n"
    "leadership sees live answers, and nothing falls through the cracks.",
    font_size=14, color=WHITE, alignment=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 4: Live Demo -- right into the demo
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.7),
    "Live Demo", font_size=36, color=ACCENT_BLUE, bold=True)
add_accent_line(slide, Inches(0.8), Inches(1.1), Inches(2), ACCENT_TEAL)

add_text_box(slide, Inches(0.8), Inches(1.4), Inches(11), Inches(0.5),
    "Pre-seeded with realistic data across the full cord blood banking lifecycle",
    font_size=16, color=LIGHT_GRAY)

# Users table
users = [
    ("Collection Team", "collection", "col123", "CBU intake, kit management, collection workflows"),
    ("Quality Assurance", "qa", "qa123", "Incident review, compliance audits, deviations"),
    ("Administrator", "admin", "admin123", "Full access, Zoho sync, user mgmt, AI config"),
]

# Table header
shape = add_shape(slide, Inches(0.8), Inches(2.2), Inches(11.7), Inches(0.5), ACCENT_BLUE)
for text, left, width in [("Role", Inches(1.0), Inches(2.5)),
    ("Username", Inches(3.8), Inches(1.5)), ("Password", Inches(5.5), Inches(1.5)),
    ("Access Scope", Inches(7.2), Inches(5.0))]:
    add_text_box(slide, left, Inches(2.22), width, Inches(0.4),
        text, font_size=13, color=WHITE, bold=True)

for i, (role, user, pwd, desc) in enumerate(users):
    top = Inches(2.8 + i * 0.65)
    bg = CARD_BG if i % 2 == 0 else RGBColor(0x15, 0x1E, 0x33)
    add_shape(slide, Inches(0.8), top, Inches(11.7), Inches(0.6), bg)
    add_text_box(slide, Inches(1.0), top + Inches(0.1), Inches(2.5), Inches(0.4),
        role, font_size=13, color=WHITE, bold=True)
    add_text_box(slide, Inches(3.8), top + Inches(0.1), Inches(1.5), Inches(0.4),
        user, font_size=12, color=ACCENT_TEAL)
    add_text_box(slide, Inches(5.5), top + Inches(0.1), Inches(1.5), Inches(0.4),
        pwd, font_size=12, color=ACCENT_TEAL)
    add_text_box(slide, Inches(7.2), top + Inches(0.1), Inches(5.0), Inches(0.4),
        desc, font_size=12, color=LIGHT_GRAY)

# Demo walkthrough
add_text_box(slide, Inches(0.8), Inches(4.9), Inches(11.7), Inches(0.4),
    "Demo Walkthrough", font_size=16, color=ACCENT_TEAL, bold=True)

demo_steps = [
    ("1", "Dashboard -- live KPIs across Mumbai, Delhi, Bangalore"),
    ("2", "CBU Tracking -- follow one unit through its full lifecycle"),
    ("3", "Storage & Temperature -- tank logs, excursions, auto-incidents"),
    ("4", "AI Copilot (Kyra) -- ask 'List recent temperature excursions'"),
    ("5", "Compliance -- incident & deviation management with audit trail"),
]

for i, (num, step) in enumerate(demo_steps):
    col = i % 3
    row = i // 3
    left = Inches(0.8 + col * 4.0)
    top = Inches(5.35 + row * 1.0)
    add_shape(slide, left, top, Inches(3.8), Inches(0.9), CARD_BG, ACCENT_TEAL)
    add_text_box(slide, left + Inches(0.15), top + Inches(0.12), Inches(3.5), Inches(0.7),
        f"{num}.  {step}", font_size=12, color=LIGHT_GRAY)

add_text_box(slide, Inches(0.8), Inches(6.55), Inches(11.7), Inches(0.6),
    "Let's jump in.", font_size=14, color=WHITE, alignment=PP_ALIGN.CENTER, bold=True)

# Save
output_path = os.path.join(os.path.dirname(__file__), "CryoSync_StemCyte_Presentation.pptx")
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")